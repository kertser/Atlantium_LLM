import logging
import re
from io import BytesIO
from pathlib import Path
from typing import Optional, Dict

from config import CONFIG
from utils.img_utils import ImageProcessor
import semchunk
from transformers import LongformerTokenizer

from openpyxl.reader.excel import load_workbook
import pymupdf  # PyMuPDF for PDFs
from PIL import Image, UnidentifiedImageError
from docx import Document

import nltk
# Download NLTK data (if not already downloaded)
nltk.download('punkt')

# Suppress MuPDF warnings
logging.getLogger("fitz").setLevel(logging.CRITICAL)
logger = logging.getLogger(__name__)

pymupdf.TOOLS.mupdf_display_errors(False)


def extract_image_safely(page, xref: int, base_image: Dict) -> Optional[Image.Image]:
    """Safely extract and convert image with multiple fallback methods."""
    image_bytes = base_image.get("image")

    if not image_bytes:
        return None

    methods = [
        lambda: Image.open(BytesIO(image_bytes)),
        lambda: Image.frombytes(
            "RGB",
            (page.get_pixmap().width, page.get_pixmap().height),
            page.get_pixmap(
                matrix=pymupdf.Matrix(1, 1),
                colorspace="rgb",
                clip=page.get_image_rects(xref)[0] if page.get_image_rects(xref) else None
            ).samples
        ),
        lambda: Image.frombytes(
            "RGB",
            (page.get_pixmap().width, page.get_pixmap().height),
            page.get_pixmap(colorspace="rgb").samples
        )
    ]

    last_error = None
    for method in methods:
        try:
            image = method()
            if image and ImageProcessor.validate_image_data(image):
                return image
        except Exception as e:
            last_error = e
            continue

    logging.debug(f"All image extraction methods failed for xref {xref}: {last_error}")
    return None

def extract_text_around_image(page, image_bbox, context_range=CONFIG.MAX_CONTEXT_RANGE):
    """Extract text around an image's location on the page."""
    try:
        blocks = page.get_text("blocks")
        image_center_y = (image_bbox[1] + image_bbox[3]) / 2
        image_center_x = (image_bbox[0] + image_bbox[2]) / 2

        nearby_text = []
        for block in blocks:
            block_center_y = (block[1] + block[3]) / 2
            block_center_x = (block[0] + block[2]) / 2

            if abs(block_center_y - image_center_y) < context_range and \
                    abs(block_center_x - image_center_x) < context_range * 2:
                text = block[4].strip()
                if text:
                    nearby_text.append(text)

        return " ".join(nearby_text)
    except Exception as e:
        logger.error(f"Error extracting text context: {e}")
        return ""

def extract_text_and_images_from_pdf(pdf_path):
    """Extracts text and images with their context from a PDF file."""
    text = ""
    image_data = []
    pdf_document = None

    try:
        pdf_document = pymupdf.open(pdf_path)
        doc_name = Path(pdf_path).name
        logger.info(f"Processing PDF document: {doc_name}")

        for page_num in range(pdf_document.page_count):
            try:
                page = pdf_document.load_page(page_num)
                page_text = page.get_text("text")
                if page_text:
                    text += page_text + "\n"

                image_list = page.get_images(full=True)

                for img_index, img in enumerate(image_list):
                    try:
                        xref = img[0]
                        base_image = pdf_document.extract_image(xref)

                        if not base_image or "image" not in base_image:
                            continue

                        try:
                            image_bytes = base_image["image"]
                            image = Image.open(BytesIO(image_bytes))

                            # Force load to verify image is valid
                            image.load()

                            # Get dimensions and calculate aspect ratio
                            width, height = image.size
                            aspect_ratio = max(width / height, height / width)

                            # Filter out small images and icons using CONFIG settings
                            if (width < CONFIG.MIN_IMAGE_SIZE or
                                    height < CONFIG.MIN_IMAGE_SIZE or
                                    max(width, height) < CONFIG.MIN_ICON_SIZE or
                                    aspect_ratio > CONFIG.MAX_ASPECT_RATIO):
                                logger.debug(
                                    f"Skipping small/icon image on page {page_num + 1}: "
                                    f"{width}x{height} pixels, aspect ratio: {aspect_ratio:.2f}"
                                )
                                continue

                            # Extract context with CONFIG.MAX_CONTEXT_RANGE
                            context = ""
                            for img_bbox in page.get_image_rects(xref):
                                context = extract_text_around_image(
                                    page,
                                    img_bbox,
                                    context_range=CONFIG.MAX_CONTEXT_RANGE
                                )
                                break

                            # Handle color modes while preserving quality
                            if image.mode in ('RGB', 'RGBA'):
                                pass  # Keep as is
                            elif image.mode == 'CMYK':
                                image = image.convert('RGB')
                            elif image.mode == 'P':
                                if 'transparency' in image.info:
                                    image = image.convert('RGBA')
                                else:
                                    image = image.convert('RGB')
                            elif image.mode in ('L', 'LA'):
                                if image.mode == 'LA':
                                    image = image.convert('RGBA')
                            else:
                                image = image.convert('RGB')

                            # Set image DPI if not already set
                            if 'dpi' not in image.info:
                                image.info['dpi'] = CONFIG.IMAGE_DPI

                            image_data.append({
                                'image': image,
                                'context': context,
                                'page_num': page_num + 1,
                                'caption': f"Image {img_index + 1} from {doc_name} (Page {page_num + 1})",
                                'dimensions': f"{width}x{height}",
                                'dpi': CONFIG.IMAGE_DPI,
                                'bits': CONFIG.IMAGE_BITS,
                                'original_mode': image.mode
                            })

                            logger.info(
                                f"Processed image {img_index + 1} from page {page_num + 1}: "
                                f"{width}x{height} pixels, mode={image.mode}"
                            )

                        except UnidentifiedImageError:
                            logger.debug(
                                f"Skipping unidentifiable image on page {page_num + 1}"
                            )
                            continue
                        except Exception as e:
                            logger.debug(
                                f"Error processing image on page {page_num + 1}: {str(e)}"
                            )
                            continue

                    except Exception as e:
                        logger.error(
                            f"Error extracting image on page {page_num + 1}: {str(e)}"
                        )
                        continue

            except Exception as e:
                logger.error(f"Error processing page {page_num + 1}: {str(e)}")
                continue

    except Exception as e:
        logger.error(f"Error processing PDF {pdf_path}: {str(e)}")
    finally:
        if pdf_document is not None:
            pdf_document.close()

    return text, image_data


def extract_text_and_images_from_word(doc_path):
    """
    Extract text and images from a Word document with enhanced image processing.

    Args:
        doc_path: Path to the Word document

    Returns:
        tuple: (extracted_text, list of image_data dictionaries)
    """
    try:
        doc = Document(doc_path)
        doc_name = Path(doc_path).name
        logger.info(f"Processing Word document: {doc_name}")

        text = "\n".join([para.text for para in doc.paragraphs])
        images_data = []

        # Extract images from relationships
        for rel in doc.part.rels.values():
            try:
                # First check if this is an image relationship
                if "image" not in rel.target_ref:  # Fixed PEP 8: E713
                    continue

                # Try to get image data, skip if not available
                if not hasattr(rel, 'target_part') or not hasattr(rel.target_part, 'blob'):
                    logger.debug(f"Skipping image in {doc_name}: No image data available")
                    continue

                image_data = rel.target_part.blob
                image = Image.open(BytesIO(image_data))

                if image.mode in CONFIG.VALID_IMAGE_MODES:
                    background = Image.new('RGB', image.size, (255, 255, 255))
                    if image.mode == 'P':
                        image = image.convert('RGBA')
                    background.paste(image, mask=image.split()[-1])
                    image = background
                elif image.mode != 'RGB':
                    image = image.convert('RGB')

                if image.width < CONFIG.MIN_IMAGE_SIZE or image.height < CONFIG.MIN_IMAGE_SIZE:
                    logger.debug(f"Skipping small image ({image.width}x{image.height}) in {doc_name}")
                    continue

                img_data = {
                    'image': image,
                    'context': "",  # Empty context as we don't extract surrounding text
                    'page_num': 1,  # Word docs don't have pages like PDFs
                    'caption': f"Image from {doc_name}",
                    'dimensions': f"{image.width}x{image.height}",
                    'format': image.format,
                    'mode': 'RGB'
                }
                images_data.append(img_data)
                logger.info(f"Processed image ({img_data['dimensions']}) from {doc_name}")

            except UnidentifiedImageError as uie:
                logger.debug(f"Invalid or corrupted image in {doc_name}: {uie}")
                continue
            except Exception as e:
                logger.debug(f"Error processing image from {doc_name}: {e}")
                continue

        logger.info(f"Completed processing {doc_name}: extracted {len(images_data)} valid images")
        return text, images_data

    except Exception as e:
        logger.error(f"Error processing Word document {doc_path}: {e}", exc_info=True)
        return "", []


def extract_text_and_images_from_excel(excel_path):
    """
    Extract text and images from an Excel file.

    Args:
        excel_path (str or Path): Path to the Excel file

    Returns:
        tuple: (extracted_text, list of image_data dictionaries)
    """
    # Suppress openpyxl warnings about headers and footers
    import warnings
    warnings.filterwarnings('ignore', category=UserWarning, message="Cannot parse header or footer*")

    workbook = None
    try:
        workbook = load_workbook(excel_path, data_only=True)
        doc_name = Path(excel_path).name
        logger.info(f"Processing Excel document: {doc_name}")

        text = ""
        images = []

        for sheet_name in workbook.sheetnames:
            try:
                worksheet = workbook[sheet_name]

                # Extract text from cells
                for row in worksheet.iter_rows(values_only=True):
                    row_text = " ".join([str(cell) if cell is not None else "" for cell in row]).strip()
                    if row_text:
                        text += row_text + "\n"

                # Extract images
                if hasattr(worksheet, '_images'):
                    for image in worksheet._images:
                        try:
                            if not hasattr(image, 'ref') or not hasattr(image.ref.image, 'content'):
                                continue

                            img = Image.open(BytesIO(image.ref.image.content))
                            width, height = img.size

                            # Skip small images
                            if width < CONFIG.MIN_IMAGE_SIZE or height < CONFIG.MIN_IMAGE_SIZE:
                                logger.debug(f"Skipping small image ({width}x{height}) in {sheet_name}")
                                continue

                            if img.mode != 'RGB':
                                img = img.convert('RGB')

                            anchor = getattr(image, 'anchor', None)
                            cell_ref = str(anchor) if anchor else "Unknown"

                            # Extract context (cells around the image)
                            context = "No context available"
                            if anchor and hasattr(anchor, '_from'):
                                try:
                                    row, col = anchor._from.row, anchor._from.col
                                    context = " ".join(
                                        str(worksheet.cell(r, c).value or "").strip()
                                        for r in range(max(1, row - 1), row + 2)
                                        for c in range(max(1, col - 1), col + 2)
                                    )
                                except Exception:
                                    logger.debug(f"Error extracting context for image in {sheet_name}")

                            images.append({
                                'image': img,
                                'context': context.strip(),
                                'page_num': 1,
                                'caption': f"Image from {doc_name} - Sheet: {sheet_name} (Cell: {cell_ref})",
                                'dimensions': f"{width}x{height}",
                                'sheet': sheet_name,
                                'cell_reference': cell_ref,
                                'source_document': str(excel_path)
                            })
                            logger.info(f"Processed image ({width}x{height}) in {sheet_name}")

                        except UnidentifiedImageError:
                            logger.debug(f"Unidentified image format in {sheet_name}")
                        except Exception as e:
                            logger.debug(f"Error processing image in {sheet_name}: {e}")

            except Exception as e:
                logger.error(f"Error processing sheet {sheet_name}: {e}")

        logger.info(f"Completed processing {doc_name}: {len(images)} images extracted")
        return text, images

    except Exception as e:
        logger.error(f"Error processing Excel file {excel_path}: {e}")
        return "", []

    finally:
        if workbook:
            try:
                workbook.close()
            except Exception as e:
                logger.debug(f"Error closing workbook: {e}")

def chunk_text_base(text: str, source_path: str, chunk_size=CONFIG.CHUNK_SIZE, overlap=CONFIG.CHUNK_OVERLAP):
    """
    Split text into chunks with overlap and enhanced metadata.
    """
    if not text or chunk_size < CONFIG.MIN_CHUNK_SIZE:
        return []

    words = text.split()
    chunks = []
    start_idx = 0
    chunk_number = 0

    while start_idx < len(words):
        end_idx = start_idx + chunk_size
        if end_idx < len(words):
            break_point = end_idx
            for i in range(max(start_idx + chunk_size - overlap, start_idx), end_idx):
                if words[i].endswith('.') or words[i].endswith('\n'):
                    break_point = i + 1
                    break
            end_idx = break_point

        chunk = {
            'text': ' '.join(words[start_idx:end_idx]),
            'metadata': {
                'source_path': source_path,
                'chunk_number': chunk_number,
                'start_idx': start_idx,
                'end_idx': end_idx,
                'relative_path': str(Path(source_path).relative_to(CONFIG.RAW_DOCUMENTS_PATH))
            }
        }
        chunks.append(chunk)
        chunk_number += 1
        start_idx = end_idx - overlap if end_idx < len(words) else end_idx

    return chunks

def chunk_text(text: str, source_path: str, chunk_size=CONFIG.CHUNK_SIZE, overlap=CONFIG.CHUNK_OVERLAP):
    """
    Split text into semantically meaningful chunks with overlap and enhanced metadata.
    Includes filtering for meaningless text such as headers, footers, disclaimers, etc.
    """
    if not text or chunk_size < CONFIG.MIN_CHUNK_SIZE:
        return []

    # Initialize the Longformer tokenizer - it can handle sequences up to 4096 tokens
    tokenizer = LongformerTokenizer.from_pretrained('allenai/longformer-base-4096')

    # Create a chunker with the specified chunk size
    chunker = semchunk.chunkerify(tokenizer, chunk_size)

    # Calculate the overlap in tokens
    overlap_tokens = int(chunk_size * overlap) if overlap < 1 else overlap

    # Filter the input text to remove irrelevant parts
    filtered_text = _filter_meaningless_text(text)

    # Obtain chunks and their offsets
    chunks, offsets = chunker(filtered_text, offsets=True, overlap=overlap_tokens)

    # Prepare the list of chunks with metadata
    chunked_data = []
    relative_path = str(Path(source_path).relative_to(CONFIG.RAW_DOCUMENTS_PATH))

    for chunk_number, (chunk_text, (start_idx, end_idx)) in enumerate(zip(chunks, offsets)):
        chunk_info = {
            'text': chunk_text,
            'metadata': {
                'source_path': source_path,
                'chunk_number': chunk_number,
                'start_idx': start_idx,
                'end_idx': end_idx,
                'relative_path': relative_path
            }
        }
        chunked_data.append(chunk_info)

    return chunked_data

def _filter_meaningless_text(text: str) -> str:
    """
    Internal function to filter out headers, footers, disclaimers, copyright lines,
    table of contents, and other irrelevant segments from the input text.
    """
    # Regular expressions to match meaningless content
    patterns = [
        # Primary TOC and dotted line patterns
        r'^.*?\.{4,}.*$',  # This will match ANY line containing 4 or more consecutive dots

        # Table of Contents headers
        r'^\s*Table\s+of\s+Contents\s*$',  # Table of Contents header
        r'^\s*Contents\s*$',  # Alternative Contents header

        # Document metadata
        r'©\s?\d{4}\s?.*?\.',  # Copyright lines
        r'^\s*Confidential\s*(?:and\s+Proprietary)?\s*$',  # Confidentiality notices
        r'^\s*All\s+[Rr]ights?\s+[Rr]eserved\s*\.?$',  # Rights reserved
        r'^\s*Draft\s+(?:version|copy)?\s*$',  # Draft markings

        # Contact information
        r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b',  # Email addresses
        r'\b(?:\+?\d{1,3}[-.\s]?)?(?:\(?\d{2,3}\)?[-.\s]?)?\d{3}[-.\s]?\d{4}\b',  # Phone numbers
        r'\bhttps?://\S+\b',  # URLs
        r'\bwww\.\S+\b',  # Web addresses

        # Document structure elements
        r'^\s*Page\s\d+\s*(?:of\s*\d+)?\s*$',  # Page indicators
        r'^\s*\d+\s*/\s*\d+\s*$',  # Page numbers like "1/10"
        r'^\s*[-_]{3,}\s*$',  # Horizontal rules
        r'^\s*[=]{3,}\s*$',  # Alternative horizontal rules

        # Headers and footers
        r'Rev(?:ision)?\.?\s?[A-Z0-9]',  # Revision indicators
        r'^\s*Last\s+(?:updated|modified|revised):\s.*$',  # Last updated lines
        r'^\s*Document\s+(?:ID|Number):\s*.*$',  # Document IDs
        r'^\s*Version:\s*\d+(?:\.\d+)*\s*$',  # Version numbers

        # Dates and timestamps
        r'\b(?:January|February|March|April|May|June|July|August|September|October|November|December)\s\d{1,2},\s\d{4}\b',
        r'\b\d{1,2}/\d{1,2}/\d{2,4}\b',  # Date formats like MM/DD/YYYY
        r'\b\d{4}-\d{2}-\d{2}\b',  # ISO date format
        r'\b(?:AM|PM|am|pm)\b',  # Time indicators

        # Business information
        r'^\s*Approved by:.*$',
        r'^\s*(?:Reviewed|Prepared) by:.*$',
        r'^\s*POB\s\d+.*$',
        r'^\s*Tel:\s?.*$',
        r'^\s*Fax:\s?.*$',

        # Disclaimers and legal text
        r'^\s*NOTICE:.*$',  # Notice blocks
        r'^\s*WARNING:.*$',  # Warning blocks
        r'^\s*Disclaimer:.*$',  # Disclaimer blocks
        r'^\s*Legal\s+Notice:.*$',  # Legal notices

        # Empty space
        r'^\s*$',  # Empty lines
        r'\n{3,}',  # Multiple consecutive newlines
    ]

    # Combine patterns and filter text
    combined_pattern = re.compile('|'.join(patterns), re.IGNORECASE | re.MULTILINE)
    filtered_text = re.sub(combined_pattern, '', text)

    # Remove excessive empty lines
    filtered_text = re.sub(r'\n{2,}', '\n', filtered_text).strip()

    return filtered_text
